from datetime import date
from io import BytesIO
from openpyxl import Workbook
from fastapi.testclient import TestClient
from app.main import app
from app.services.excel import inspect_excel, normalize_excel
from app.services.database import authenticated, AuthToken
from app.routes import uploads


def fixture():
    wb = Workbook()
    wb.active.title = 'Summary'
    wb.active.append(['Laporan penjualan'])
    ws = wb.create_sheet('Detail')
    ws.merge_cells('A1:E1')
    ws['A1'] = 'PT Contoh — Laporan Penjualan'
    ws.append(['Periode Oktober 2026'])
    ws.append([])
    ws.append(['DATE', 'PRODUCT', 'QTY', 'UNIT PRICE', 'UNIT COST'])
    ws.append([date(2026, 10, 1), 'Lensa', 2, 120000, 80000])
    ws.append([date(2026, 10, 2), 'Frame', 1, 90000, 50000])
    wb.create_sheet('Kosong')
    buf = BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def test_multi_sheet_detection_and_original_preview():
    content = fixture()
    initial = inspect_excel(content)
    assert initial['sheet'] == ''
    assert initial['sheets'] == ['Summary', 'Detail', 'Kosong']
    result = inspect_excel(content, 'Detail')
    assert result['recognized'] is True
    assert result['header_row'] == 4
    assert result['sample'][0][1:] == ['Lensa', '2', '120000', '80000']
    records = normalize_excel(content, 'Detail', result['header_row'], result['mapping'])
    assert len(records) == 2
    assert records[0]['price'] == '120000'
    for sheet in ['Summary', 'Kosong']:
        assert inspect_excel(content, sheet)['recognized'] is False


def test_detection_not_fixed_to_fourth_row():
    wb = Workbook()
    for _ in range(13):
        wb.active.append(['Keterangan laporan'])
    wb.active.append(['STORE', 'SKU_DESC', 'QTY_SALES', 'SALES_VALUE'])
    wb.active.append(['Bali', 'Produk A', 2, 20000])
    buf = BytesIO()
    wb.save(buf)
    result = inspect_excel(buf.getvalue())
    assert result['header_row'] == 14
    assert 'price' not in result['mapping']


def test_readable_pivot_is_not_a_complete_transaction_mapping():
    """A readable table and a mapped quantity do not provide missing financial fields."""
    import pytest
    from fastapi import HTTPException
    wb = Workbook()
    wb.active.append(['SKU_DESC', 'REGION', 'Sum of QTY_SALES'])
    wb.active.append(['PAKET PAYDAY', 'Makassar', 8])
    wb.active.append(['PAKET PAYDAY', 'Bali', 41])
    wb.active.append(['Grand Total', None, 49])
    buf = BytesIO()
    wb.save(buf)
    wb.close()
    result = inspect_excel(buf.getvalue())
    assert result['recognized'] is True  # Readability only; not permission to save.
    assert result['mapping']['product'] == 0
    assert result['mapping']['quantity'] == 2
    for field in ['date', 'price', 'unit_cost']:
        assert field not in result['mapping']
    with pytest.raises(HTTPException) as error:
        normalize_excel(buf.getvalue(), result['sheet'], result['header_row'], result['mapping'])
    assert error.value.status_code == 400


def test_no_database_write_before_explicit_commit(monkeypatch):
    import json
    calls = []
    token = AuthToken('test')
    token.user_id = 'user-a'
    app.dependency_overrides[authenticated] = lambda: token
    async def database(*args, **kwargs):
        calls.append(kwargs['body'])
        return []
    monkeypatch.setattr(uploads, 'database', database)
    try:
        with TestClient(app) as client:
            content = fixture()
            file = {'file': ('test.xlsx', content)}
            profile = client.post('/api/import/inspect', files=file, data={'sheet': 'Detail'}).json()
            data = {'sheet': 'Detail', 'header_row': profile['header_row'], 'mapping': json.dumps(profile['mapping'])}
            assert client.post('/api/import/preview', files=file, data={**data, 'confirm': 'true'}).status_code == 200
            assert client.post('/api/import/commit', files=file, data=data).status_code == 200
            assert calls == []
            assert client.post('/api/import/commit', files=file, data={**data, 'confirm': 'true'}).json() == {'imported': 2}
            assert len(calls) == 1
    finally:
        app.dependency_overrides.clear()


def test_duplicate_columns_explain_which_fields_conflict_and_never_write(monkeypatch):
    import json
    calls = []
    token = AuthToken('test')
    token.user_id = 'duplicate-column-user'
    async def database(*args, **kwargs):
        calls.append(kwargs)
        return []
    monkeypatch.setattr(uploads, 'database', database)
    app.dependency_overrides[authenticated] = lambda: token
    try:
        with TestClient(app) as client:
            content = fixture()
            profile = inspect_excel(content, 'Detail')
            for field, column, labels in [
                ('unit_cost', 3, ['Harga jual aktual per unit', 'HPP per unit']),
                ('category', 1, ['Produk', 'Kategori']),
                ('date', 2, ['Tanggal transaksi', 'Jumlah unit']),
            ]:
                mapping = {**profile['mapping'], field: column}
                data = {'sheet': 'Detail', 'header_row': profile['header_row'],
                        'mapping': json.dumps(mapping), 'confirm': 'true'}
                for endpoint in ['preview', 'commit']:
                    response = client.post('/api/import/'+endpoint, files={'file': ('source.xlsx', content)}, data=data)
                    assert response.status_code == 400
                    assert all(label in response.json()['detail'] for label in labels)
                    assert 'file yang lebih lengkap' in response.json()['detail']
            assert calls == []
    finally:
        app.dependency_overrides.clear()


def test_full_http_flow_from_empty_data_to_reports(monkeypatch):
    """Exercises actual FastAPI endpoints, Excel parser and financial calculations."""
    import json
    from uuid import uuid4
    from app.routes import dashboard
    from openpyxl import load_workbook
    state = {'refraksi_sales': [], 'refraksi_expenses': []}
    token = AuthToken('test')
    token.user_id = 'flow-user'
    async def all_records(token, table):
        return list(state[table])
    async def database(token, method, table, *, body=None, params=None, **kwargs):
        rows = state[table]
        if method == 'POST':
            values = body if isinstance(body, list) else [body]
            result = [{**row, 'id': row.get('id', str(uuid4()))} for row in values]
            rows.extend(result)
            return result
        selected = [row for row in rows if row['id'] == params['id'][3:]]
        if method == 'PATCH':
            for row in selected:
                row.update(body)
        elif method == 'DELETE':
            state[table] = [row for row in rows if row not in selected]
        return selected
    app.dependency_overrides[authenticated] = lambda: token
    monkeypatch.setattr(dashboard, 'database', database)
    monkeypatch.setattr(dashboard, 'all_records', all_records)
    monkeypatch.setattr(uploads, 'database', database)
    try:
        with TestClient(app) as client:
            assert client.get('/api/dashboard?year=2026').json()['sales'] == []
            assert client.get('/api/analysis?year=2026').json()['has_sales'] is False
            files = {'file': ('source.xlsx', fixture())}
            profile = client.post('/api/import/inspect', files=files).json()
            assert profile['sheet'] == ''
            profile = client.post('/api/import/inspect', files=files, data={'sheet': 'Detail'}).json()
            data = {'sheet': 'Detail', 'header_row': profile['header_row'], 'mapping': json.dumps(profile['mapping'])}
            preview = client.post('/api/import/preview', files=files, data=data).json()
            assert preview['count'] == 2
            assert state['refraksi_sales'] == []
            # Cancelling needs no server request; inspect/preview have no writes.
            assert client.post('/api/import/commit', files=files, data={**data, 'confirm': 'true'}).json()['imported'] == 2
            report = client.get('/api/dashboard?year=2026&month=10').json()
            assert report['metrics']['sales'] == '330000'
            assert report['metrics']['gross'] == '120000'
            analysis = client.get('/api/analysis?year=2026&month=10').json()
            assert analysis['has_sales'] is True
            assert analysis['leaders']['sales']['product'] == 'Lensa'
            assert analysis['comparisons']['sales']['growth'] is None
            assert client.get('/api/analysis?year=2026&month=9').json()['has_sales'] is False
            assert len(client.get('/api/analysis?year=2026&product=Frame').json()['products']) == 1
            sale = dict(date='2026-10-03', product='CRUD', quantity=1, price='100', unit_cost='50')
            added = client.post('/api/sales', json=sale)
            assert added.status_code == 200
            record_id = added.json()[0]['id']
            assert client.put('/api/sales/'+record_id, json={**sale, 'quantity': 2}).status_code == 200
            expense = client.post('/api/expenses', json=dict(date='2026-10-03', description='Sewa', amount='10000'))
            assert expense.status_code == 200
            assert client.get('/api/dashboard?year=2026&month=10').json()['metrics']['net'] == '110100'
            assert client.delete('/api/sales/'+record_id).status_code == 200
            exported = client.get('/api/export?year=2026&product=Frame')
            wb = load_workbook(BytesIO(exported.content))
            assert wb['Sales'].max_row == 2
            assert wb['Sales']['D2'].value == 'Frame'  # id,date,customer,product
            wb.close()
    finally:
        app.dependency_overrides.clear()
