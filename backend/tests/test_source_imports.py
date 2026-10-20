from io import BytesIO
from datetime import date
from decimal import Decimal
from uuid import uuid4
import pytest
from openpyxl import Workbook, load_workbook
from fastapi import HTTPException
from fastapi.testclient import TestClient
from app.services.source_excel import inspect_sources, prepare_source
from app.services.sources import recap_sales, overlaps
from app.services.finance import report
from app.services.analysis import analyse
from app.services.excel import export_excel
from app.models.records import Filters
from app.main import app


def workbook(name, header, rows, header_row=1, title=None):
    wb = Workbook()
    ws = wb.active
    ws.title = name
    if title:
        ws['A1'] = title
    for i, label in enumerate(header, 1):
        ws.cell(header_row, i, label)
    for r, values in enumerate(rows, header_row + 1):
        for c, value in enumerate(values, 1):
            if value is not None:
                ws.cell(r, c, value)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def preview(content, name, **options):
    return prepare_source(content, name, {'year': 2026, 'cost_basis': 'none', **options}, 'account-a', 'source.xlsx')


def test_master_uses_sku_not_create_date_and_never_adds_sales():
    content = workbook('master product', ['SKU FINAL', 'UPC', 'DESCRIPTION FINAL', 'DEPT FINAL', 'Create Date', 'SKU FINAL'],
                       [['001', '123', 'Produk', 'Food', date(2020, 1, 1), '001']], 6, 'MASTER BARANG')
    profile = inspect_sources(content)['source_profile']
    assert profile['kind'] == 'master'
    view, payload = preview(content, 'master product')
    assert view['can_import']
    assert payload['records'][0]['sku'] == '001'
    assert 'date' not in payload['records'][0]
    assert recap_sales([payload]) == []
    assert view['totals']['sales'] is None


def test_monthly_uses_all_sales_not_promo_and_collapses_only_identical_rows():
    header = ['SKU', 'SKU_DESC', 'BULAN', 'STORE', 'SALES QTY PROMO', 'VALUE PROMO', 'RAFAKSI',
              'SALES QTY (ALL STORE)(1 BULAN)', 'SALES VALUE (ALL STORE)(1 BULAN)', 'COST + PPn']
    row = ['001', 'Produk', 'JANUARI', 'BALI', 3, 15000, 15000, 7, 100000, 100.1234]
    content = workbook('JANUARI', header, [row, row], 4, 'PERIODE JANUARI 2026')
    view, payload = preview(content, 'JANUARI', cost_basis='tax_cost')
    assert view['can_import']
    assert view['count'] == 1
    assert view['excluded']['duplicate'] == 1
    assert Decimal(view['totals']['sales']) == 100000
    assert Decimal(view['totals']['hpp']) == Decimal('700.8638')
    assert payload['records'][0]['period_end'] == '2026-01-31'
    different_program = list(row)
    different_program[3] = 'NASIONAL'
    same_scope, _ = preview(workbook('JANUARI', header, [row, different_program], 4), 'JANUARI')
    assert same_scope['count'] == 1
    assert same_scope['excluded']['duplicate'] == 1
    changed = [*row[:8], 90000, row[-1]]
    bad, _ = preview(workbook('JANUARI', header, [row, changed], 4), 'JANUARI')
    assert not bad['can_import']
    assert any('nilai berbeda' in item['message'] for item in bad['issues'])
    with pytest.raises(HTTPException):
        preview(content, 'JANUARI', year=2025)


def test_payday_requires_real_period_and_preserves_store_identity():
    content = workbook('DETAIL SALES', ['STORE', 'STORE_NAME', 'SKU', 'SKU_DESC', 'QTY_SALES', 'SALES_VALUE', 'REGION'],
                       [[1, 'Toko A', '001', 'Paket', 3, 100, 'Bali'], [2, 'Toko B', '001', 'Paket', 2, 80, 'Bali']])
    with pytest.raises(HTTPException):
        preview(content, 'DETAIL SALES')
    view, payload = preview(content, 'DETAIL SALES', start='2026-09-25', end='2026-09-30')
    assert view['can_import']
    assert view['count'] == 2
    assert view['totals']['sales'] == '180'
    assert view['totals']['hpp'] is None
    assert {r['location'] for r in payload['records']} == {'Toko A', 'Toko B'}
    sales = recap_sales([payload])
    results = report(sales, [], Filters(year=2026))
    assert results['metrics']['sales'] == '180'
    assert results['metrics']['gross'] is None
    assert all(row['gross'] is None for row in results['timeline'] if row['month'] == '2026-09')
    assert results['transactions'] == []
    assert report(sales, [], Filters(start=date(2026, 9, 26)))['sales'] == []
    analysis = analyse(sales, [], Filters(year=2026))
    assert analysis['comparisons']['gross']['current'] is None
    assert analysis['leaders']['gross'] is None
    wb = load_workbook(export_excel(results))
    assert wb['Sales'].max_row == 1
    assert wb['Recaps'].max_row == 3


def test_supplier_claims_do_not_become_sales_or_expenses():
    content = workbook('DETAIL RAFRAKSI FEBRUARI', ['CODE SUPPLIER', 'SUPPLIER NAME', 'Sum of SALES QTY PROMO', 'Sum of RAFAKSI2'],
                       [['S1', 'Supplier A', 12, 34000], ['Grand Total', None, 12, 34000]], 6)
    view, payload = preview(content, 'DETAIL RAFRAKSI FEBRUARI')
    assert view['can_import']
    assert view['totals']['claim'] == '34000'
    assert view['totals']['sales'] is None
    assert view['excluded']['total'] == 1
    assert recap_sales([payload]) == []
    assert payload['records'][0]['period_end'] == '2026-02-28'


def test_wide_qsr_does_not_import_average_or_total_and_exposes_missing_periods():
    header = ['NO', 'SKU', 'SKU DESC', 'JANUARI', 'FEBRUARI', 'MARET', 'AVG PER DAY', 'APRIL', 'MEI', 'JUNI', 'JULI', 'AGUSTUS', 'TOTAL',
              'JANUARI', 'FEBRUARI', 'MARET', 'APRIL', 'MEI', 'JUNI', 'JULI', 'AGUSTUS', 'TOTAL']
    row = [1, '001', 'Produk', 3, 0, 0, 999, 0, 0, 0, 0, 0, 1002, 100, 'X', 0, 0, 0, 0, 0, 0, 100]
    content = workbook('SALES QSR', header, [row], 2)
    view, payload = preview(content, 'SALES QSR')
    assert view['can_import']
    assert view['count'] == 1
    assert payload['records'][0]['quantity'] == 3
    assert view['totals']['sales'] == '100'
    assert view['excluded']['unavailable_period'] == 1
    assert view['unavailable_periods'][0]['period_start'] == '2026-02-01'


def test_event_preserves_aggregate_precision_and_no_fake_unit_price():
    content = workbook('DOUBLE DATE OKTOBER', ['SKU', 'SKU_DESC', 'REGION', 'AREA', 'QTY SALES (10 OKTOBER)', 'SALES VALUE (10 OKTOBER)', 'COST', 'COST+PPN'],
                       [['001', 'Produk', 'BALI', 'Airport', 3, 100, 1, 1.1234]], 4)
    view, payload = preview(content, 'DOUBLE DATE OKTOBER', cost_basis='tax_cost')
    assert view['can_import']
    assert payload['records'][0]['period_start'] == payload['records'][0]['period_end'] == '2026-10-10'
    assert payload['records'][0]['total_hpp'] == '3.3702'
    assert 'price' not in payload['records'][0]
    other = prepare_source(content, 'DOUBLE DATE OKTOBER', dict(year=2026, cost_basis='tax_cost'), 'account-a', 'renamed.xlsx')[1]
    assert other['fingerprint'] == payload['fingerprint']
    assert other['id'] == payload['id']


def test_overlap_detects_aggregate_and_legacy_transactions_but_not_other_stores():
    a = dict(sku='1', product='Produk', period_start='2026-01-01', period_end='2026-01-31', location='Bali', aggregate_location=True)
    assert overlaps(a, dict(product='Produk', period_start='2026-01-10', period_end='2026-01-10', location='Toko A'))
    assert not overlaps(a, {**a, 'period_start':'2026-02-01', 'period_end':'2026-02-28'})
    assert not overlaps({**a, 'aggregate_location':False}, {**a, 'location':'Toko B', 'aggregate_location':False})


def test_commit_requires_latest_preview_explicit_confirm_and_no_overlap(monkeypatch):
    from app.services import sources
    from app.services.database import authenticated, AuthToken
    from app.routes import dashboard, uploads
    token = AuthToken('test')
    token.user_id = 'account-a'
    state = {'refraksi_sales': [], 'refraksi_expenses': [], 'refraksi_sources': []}
    async def read(token, table):
        return state[table]
    async def write(token, method, table, *, body=None, params=None):
        if method == 'POST':
            state[table].append(body)
            return [body]
        return []
    for module in [sources, dashboard]:
        monkeypatch.setattr(module, 'all_records', read)
    monkeypatch.setattr(uploads, 'database', write)
    app.dependency_overrides[authenticated] = lambda: token
    try:
        client = TestClient(app)
        content = workbook('DETAIL SALES', ['SKU', 'SKU_DESC', 'QTY_SALES', 'SALES_VALUE'], [['1', 'Produk', 3, 100]])
        files = {'file': ('source.xlsx', content)}
        import json
        data = {'sheet':'DETAIL SALES', 'options':json.dumps(dict(year=2026, start='2026-01-01', end='2026-01-31'))}
        view = client.post('/api/import/sources/preview', files=files, data=data).json()
        assert view['can_import']
        assert not state['refraksi_sources']
        assert client.post('/api/import/sources/commit', files=files, data={**data, 'confirm':'true'}).status_code == 400
        assert not state['refraksi_sources']
        assert client.post('/api/import/sources/commit', files=files, data={**data, 'confirm':'true','preview_signature':view['preview_signature']}).json()['imported'] == 1
        assert client.get('/api/dashboard?year=2026').json()['metrics']['sales'] == '100'
        assert client.get('/api/analysis?year=2026').json()['comparisons']['gross']['current'] is None
        again = client.post('/api/import/sources/preview', files=files, data=data).json()
        assert not again['can_import']
        assert any('sudah tersimpan' in item['message'] for item in again['issues'])
        assert len(state['refraksi_sources']) == 1
    finally:
        app.dependency_overrides.clear()
    assert TestClient(app).post('/api/import/sources/preview', files=files, data=data).status_code in (401,503)


def test_cached_formula_requires_acknowledgement_and_uses_decimal_xml_exactly():
    from zipfile import ZipFile
    content = workbook('DETAIL SALES', ['SKU','SKU_DESC','QTY_SALES','SALES_VALUE','COST+PPN'], [['1','Produk',3,100,'=1.1234']])
    output = BytesIO()
    with ZipFile(BytesIO(content)) as original, ZipFile(output, 'w') as edited:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename == 'xl/worksheets/sheet1.xml':
                data = data.replace(b'<f>1.1234</f><v />', b'<f>1.1234</f><v>1.123456789012345</v>').replace(b'<f>1.1234</f><v></v>', b'<f>1.1234</f><v>1.123456789012345</v>')
            edited.writestr(info, data)
    args = dict(start='2026-01-01',end='2026-01-31',cost_basis='tax_cost')
    view, _ = preview(output.getvalue(), 'DETAIL SALES', **args)
    assert not view['can_import']
    assert view['cached_formula_count'] == 1
    view, payload = preview(output.getvalue(), 'DETAIL SALES', accept_cached=True, **args)
    assert view['can_import']
    assert payload['records'][0]['total_hpp'] == '3.370370367037035'


def test_missing_schema_falls_back_to_existing_dashboard_without_masking_outages(monkeypatch):
    import asyncio
    from app.services import sources
    async def missing(*args):
        raise sources.SourceSchemaMissing()
    monkeypatch.setattr(sources, 'all_records', missing)
    assert asyncio.run(sources.source_datasets('test')) == ([], False)
    async def outage(*args):
        raise HTTPException(503,'Network unavailable')
    monkeypatch.setattr(sources, 'all_records', outage)
    with pytest.raises(HTTPException):
        asyncio.run(sources.source_datasets('test'))


def test_optional_master_reference_error_is_unavailable_not_zero():
    content = workbook('master product', ['SKU FINAL', 'UPC', 'DESCRIPTION FINAL', 'DEPT FINAL'], [['001', '#N/A', 'Produk', 'Food']])
    view, payload = preview(content, 'master product')
    assert view['can_import']
    assert view['excluded']['optional_unavailable'] == 1
    assert payload['records'][0]['upc'] == ''
    bad = workbook('master product', ['SKU FINAL', 'UPC', 'DESCRIPTION FINAL', 'DEPT FINAL'], [['001', '123', '#N/A', 'Food']])
    view, _ = preview(bad, 'master product')
    assert not view['can_import']
