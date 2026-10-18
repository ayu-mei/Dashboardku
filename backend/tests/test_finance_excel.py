from datetime import date
from decimal import Decimal
from io import BytesIO
import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from app.models.records import Filters
from app.services.finance import report
from app.services.excel import inspect_excel, normalize_excel, export_excel


def sales():
    return [
        dict(
            id="a",
            date="2026-01-10",
            product="Lensa",
            customer="Ani",
            quantity=2,
            price="500000",
            unit_cost="300000",
            location="Bali",
            category="Optik",
        ),
        dict(
            id="b",
            date="2026-02-10",
            product="Frame",
            customer="Budi",
            quantity=1,
            price="200000",
            unit_cost="100000",
            location="Jakarta",
            category="Frame",
        ),
    ]


def test_finance_weighted_margin_and_filters():
    expenses = [dict(id="e", date="2026-01-12", description="Sewa", amount="100000")]
    january = report(sales(), expenses, Filters(year=2026, month=1))
    assert january["metrics"] == dict(
        sales="1000000",
        hpp="600000",
        gross="400000",
        gross_margin="40.0",
        operational="100000",
        net="300000",
        net_margin="30.0",
    )
    combined = report(sales(), expenses, Filters(year=2026))
    assert (
        Decimal(combined["metrics"]["gross_margin"])
        == Decimal(500000) / Decimal(1200000) * 100
    )
    selected = report(sales(), expenses, Filters(year=2026, product="Frame"))
    assert selected["metrics"]["sales"] == "200000"
    assert (
        selected["metrics"]["net"] == "0"
    )  # All period costs remain, no assumed allocation.
    assert (
        report(
            sales(), expenses, Filters(start=date(2026, 2, 1), end=date(2026, 2, 28))
        )["metrics"]["sales"]
        == "200000"
    )
    assert (
        report(sales(), expenses, Filters(location="Bali"))["metrics"]["sales"]
        == "1000000"
    )
    assert (
        report(sales(), expenses, Filters(category="Frame"))["metrics"]["sales"]
        == "200000"
    )
    empty = report([], expenses, Filters(year=2026))
    assert empty["metrics"]["gross_margin"] is None
    assert empty["metrics"]["net_margin"] is None
    assert empty["metrics"]["net"] == "-100000"
    exact = report(
        [
            dict(
                date="2026-01-01",
                product="Tiny",
                quantity=3,
                price="0.10",
                unit_cost="0.03",
            )
        ],
        [],
        Filters(),
    )
    assert exact["metrics"]["gross"] == "0.21"


def workbook(headers=None, row=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Detail"
    ws.append(
        headers or ["Tanggal Asli", "SKU_DESC", "QTY_SALES", "Harga Aktual", "COST"]
    )
    ws.append(row or [date(2026, 1, 10), "Lensa", 2, 500000, 300000])
    buffer = BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()


MAPPING = {"date": 0, "product": 1, "quantity": 2, "price": 3, "unit_cost": 4}


def test_inspection_and_mapping_preserve_actual_names():
    content = workbook()
    result = inspect_excel(content)
    assert [col["name"] for col in result["columns"]] == [
        "Tanggal Asli",
        "SKU_DESC",
        "QTY_SALES",
        "Harga Aktual",
        "COST",
    ]
    records = normalize_excel(content, "Detail", 1, MAPPING, "user-a")
    assert records[0]["price"] == "500000"
    assert records[0]["date"] == "2026-01-10"
    assert (
        normalize_excel(content, "Detail", 1, MAPPING, "user-a")[0]["id"]
        == records[0]["id"]
    )
    assert (
        normalize_excel(content, "Detail", 1, MAPPING, "user-b")[0]["id"]
        != records[0]["id"]
    )


def test_invalid_uploads_and_subtotals():
    for content in [b"", b"not an xlsx"]:
        with pytest.raises(HTTPException):
            inspect_excel(content)
    with pytest.raises(HTTPException):
        normalize_excel(workbook(), "Detail", 1, {"product": 1})
    with pytest.raises(HTTPException):
        normalize_excel(
            workbook(row=[date(2026, 1, 10), "Grand Total", 2, 500000, 300000]),
            "Detail",
            1,
            MAPPING,
        )
    with pytest.raises(HTTPException):
        normalize_excel(
            workbook(row=[date(2026, 1, 10), "Lensa", 2, "=1+1", 300000]),
            "Detail",
            1,
            MAPPING,
        )
    with pytest.raises(HTTPException):
        normalize_excel(
            workbook(row=["Januari", "Lensa", 2, 500000, 300000]), "Detail", 1, MAPPING
        )


def test_excel_export_is_not_formula_injection():
    data = report(
        [
            dict(
                id="a",
                date="2026-01-10",
                product="=1+1",
                quantity=2,
                price="500000",
                unit_cost="300000",
            )
        ],
        [],
        Filters(),
    )
    wb = load_workbook(export_excel(data), data_only=False)
    assert wb.sheetnames == ["Sales", "Expenses", "Summary"]
    assert wb["Sales"]["D2"].value == "=1+1"
    assert wb["Sales"]["D2"].data_type == "s"
    wb.close()
