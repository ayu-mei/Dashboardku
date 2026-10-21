from datetime import datetime, date
from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from openpyxl.utils.exceptions import InvalidFileException
from uuid import uuid5, NAMESPACE_URL
import hashlib
import json
import re
import pandas as pd
from openpyxl import load_workbook, Workbook
from fastapi import HTTPException
from app.models.records import Sale
from app.utils.config import settings

MAX_ROWS = 20000


def checked_workbook(content):
    try:
        with ZipFile(BytesIO(content)) as archive:
            # XLSX is compressed: legitimate multi-sheet workbooks can expand
            # well beyond the upload size. Keep a separate, bounded limit.
            limit = settings().max_workbook_bytes
            if sum(info.file_size for info in archive.infolist()) > limit:
                raise HTTPException(
                    400, f"Workbook berisi terlalu banyak data setelah dibuka (maksimal {limit // (1024 * 1024)} MB). "
                    "Simpan hanya sheet yang diperlukan sebagai file Excel terpisah, lalu coba lagi."
                )
            if any(
                "vbaproject" in info.filename.lower() for info in archive.infolist()
            ):
                raise HTTPException(400, "Excel dengan macro tidak diterima.")
        return load_workbook(
            BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except (BadZipFile, KeyError, ValueError, OSError, ParseError, InvalidFileException):
        raise HTTPException(
            400, "File bukan Excel .xlsx yang valid atau file kosong/rusak."
        )


# Only unambiguous unit-price/cost names are suggested. SALES_VALUE is a total.
ALIASES = {
    "date": {"date", "tanggal", "tanggal transaksi", "transaction date", "tanggal asli", "sales date", "tgl", "tgl transaksi"},
    "product": {"product", "produk", "sku desc", "sku description", "product name", "nama produk", "nama barang", "nama item", "item desc"},
    "quantity": {"quantity", "qty", "qty sales", "qty sold", "jumlah", "jumlah unit", "qty penjualan", "sales qty", "jumlah penjualan", "sum of qty sales", "sum of qty", "sum of quantity"},
    "price": {"price", "unit price", "harga aktual", "harga jual aktual per unit", "harga satuan", "harga jual satuan", "harga jual per unit", "actual unit price"},
    "unit_cost": {"unit cost", "hpp per unit", "hpp satuan", "cost"},
    "customer": {"customer", "pelanggan", "nama pelanggan"},
    "category": {"category", "kategori", "department"},
    "location": {"location", "lokasi", "store name", "cabang"},
}
STRUCTURE_NAMES = set().union(*ALIASES.values()) | {
    "store", "sku", "sales", "sales value", "store name", "cost ppn", "region", "area"
}
FIELD_LABELS = {
    "date": "Tanggal transaksi", "product": "Produk", "quantity": "Jumlah unit",
    "price": "Harga jual aktual per unit", "unit_cost": "HPP per unit",
    "customer": "Pelanggan", "category": "Kategori", "location": "Cabang / lokasi",
}


def column_name(value):
    return re.sub(r"[^a-z0-9]+", " ", str(value).lower()).strip()


def detect_header(rows):
    candidates = []
    for index, row in enumerate(rows[:50]):
        occupied = [i for i, value in enumerate(row) if value is not None and str(value).strip()]
        texts = [i for i in occupied if isinstance(row[i], str) and not row[i].startswith("=")]
        if len(texts) < 2 or len(texts) / max(len(occupied), 1) < .75:
            continue
        names = [column_name(row[i]) for i in texts]
        if len(set(names)) < len(names) * .7:
            continue
        following = [r for r in rows[index + 1:index + 7] if any(v is not None for v in r)]
        if not following:
            continue
        coverage = [sum(i < len(r) and r[i] is not None for i in occupied) / len(occupied) for r in following]
        consistent = sum(value >= .6 for value in coverage) / len(coverage)
        recognized = sum(name in STRUCTURE_NAMES for name in names)
        typed = any(any(isinstance(v, (int, float, date, datetime)) for v in r) for r in following)
        if consistent < .6 or (recognized < 2 and not (typed and len(texts) >= 3)):
            continue
        candidates.append((recognized * 3 + consistent * 2 + min(len(texts), 15) * .1, -index))
    return -max(candidates)[1] + 1 if candidates else None


def inspect_excel(content, sheet=None, header_row=None):
    wb = checked_workbook(content)
    try:
        if sheet and sheet not in wb.sheetnames:
            raise HTTPException(400, "Sheet tidak ditemukan.")
        # Multiple sheets require an explicit choice; never combine sheets.
        if not sheet and len(wb.sheetnames) > 1:
            return {"sheets": wb.sheetnames, "sheet": "", "recognized": False, "columns": [], "sample": []}
        ws = wb[sheet or wb.sheetnames[0]]
        initial = list(ws.iter_rows(max_row=60, values_only=True))
        header_row = header_row or detect_header(initial)
        if not header_row:
            return {"sheets": wb.sheetnames, "sheet": ws.title, "recognized": False,
                    "columns": [], "sample": [],
                    "message": "Struktur data belum dapat dikenali. Pilih sheet lain atau periksa format file Excel."}
        rows = list(ws.iter_rows(min_row=header_row, max_row=header_row + 10, values_only=True))
        columns = [{"index": i, "name": str(value).strip()} for i, value in enumerate(rows[0])
                   if value is not None and str(value).strip()]
        mapping = {}
        for field, aliases in ALIASES.items():
            matches = [col["index"] for col in columns if column_name(col["name"]) in aliases]
            if len(matches) == 1:
                mapping[field] = matches[0]
        from app.services.excel_reports import describe_report
        report = describe_report(ws, header_row)
        return {
            "sheets": wb.sheetnames, "sheet": ws.title, "header_row": header_row,
            "recognized": True, "columns": columns, "mapping": mapping,
            "report": report,
            "sample": [[str(v) if v is not None else "" for v in row] for row in rows[1:] if any(v is not None for v in row)],
            "warning": "Pilih sheet detail, bukan subtotal/pivot. Tanggal, produk, jumlah, harga jual aktual per unit, dan HPP per unit wajib tersedia. SALES_VALUE adalah total, bukan harga satuan. Periksa basis pajak COST sebelum mengimpor.",
        }
    finally:
        wb.close()


def normalize_excel(content, sheet, header_row, mapping, owner="local"):
    required = ["date", "product", "quantity", "price", "unit_cost"]
    if not isinstance(mapping, dict) or any(
        mapping.get(field) is None for field in required
    ):
        raise HTTPException(
            400,
            "Pemetaan belum lengkap. Wajib: tanggal, produk, jumlah, harga jual aktual per unit, dan HPP satuan. Data agregat tanpa field ini belum dapat diimpor.",
        )
    indices = [value for value in mapping.values() if value is not None]
    if any(type(i) is not int or i < 0 for i in indices) or any(key not in FIELD_LABELS for key in mapping):
        raise HTTPException(400, "Pilihan kolom tidak valid. Pilih kembali kolom dari daftar.")
    uses = {}
    for field, index in mapping.items():
        if index is not None:
            uses.setdefault(index, []).append(FIELD_LABELS[field])
    conflicts = [" dan ".join(labels) for labels in uses.values() if len(labels) > 1]
    if conflicts:
        raise HTTPException(
            400,
            "Kolom Excel yang sama dipakai untuk: " + "; ".join(conflicts)
            + ". Pilih kolom berbeda untuk setiap isian, atau kosongkan isian opsional. "
              "Jika datanya tidak tersedia, gunakan sheet atau file yang lebih lengkap.",
        )
    wb = checked_workbook(content)
    try:
        if sheet not in wb.sheetnames:
            raise HTTPException(400, "Sheet tidak ditemukan.")
        ws = wb[sheet]
        rows = []
        for row in ws.iter_rows(min_row=header_row + 1):
            if not any(cell.value is not None for cell in row):
                continue
            if len(rows) >= MAX_ROWS:
                raise HTTPException(400, "Maksimal 20.000 baris per impor.")
            number = row[0].row
            values = {}
            for field, index in mapping.items():
                if index is None:
                    continue
                if index >= len(row):
                    raise HTTPException(
                        400, f"Kolom tidak ditemukan pada baris {number}."
                    )
                cell = row[index]
                if cell.data_type == "f":
                    raise HTTPException(
                        400,
                        f"Baris {number} berisi formula pada kolom yang dipilih. Gunakan salinan ekspor nilai transaksi aktual.",
                    )
                value = cell.value
                if field == "date":
                    if isinstance(value, datetime):
                        value = value.date()
                    if not isinstance(value, date):
                        try:
                            value = date.fromisoformat(str(value))
                        except ValueError:
                            raise HTTPException(
                                400,
                                f"Baris {number}: tanggal harus tanggal Excel atau YYYY-MM-DD, bukan nama bulan/periode.",
                            )
                values[field] = value
            try:
                record = Sale.model_validate(values).model_dump(mode="json")
                if record["product"].strip().casefold() in {
                    "total",
                    "grand total",
                    "subtotal",
                    "sub total",
                }:
                    raise ValueError("Subtotal is not a transaction")
            except Exception:
                raise HTTPException(
                    400,
                    f"Baris {number} tidak valid. Periksa tanggal, jumlah positif bulat, teks, dan uang non-negatif maksimal 2 desimal. Jangan pilih subtotal/Grand Total.",
                )
            rows.append(record)
        if not rows:
            raise HTTPException(
                400, "Tidak ada baris transaksi pada sheet yang dipilih."
            )
        # pandas cleaning preserves numeric decimal strings; no fillna(0), rounding, or silent row deletion.
        records = (
            pd.DataFrame(rows)
            .where(pd.notna(pd.DataFrame(rows)), None)
            .to_dict(orient="records")
        )
        signature = hashlib.sha256(
            content
            + json.dumps(
                {"sheet": sheet, "header": header_row, "mapping": mapping},
                sort_keys=True,
            ).encode()
        ).hexdigest()
        for index, record in enumerate(records):
            record["id"] = str(
                uuid5(NAMESPACE_URL, owner + ":" + signature + ":" + str(index))
            )
        return records
    finally:
        wb.close()


def export_excel(data):
    data = {**data, "metrics": {key: value for key, value in data["metrics"].items() if key not in {"operational", "net", "net_margin"}}, "note": "Angka mengikuti data sumber. Laba kotor = penjualan dikurangi HPP; HPP yang belum tersedia tidak dianggap nol."}
    wb = Workbook()
    wb.remove(wb.active)
    groups = [
        (
            "Sales",
            data.get("transactions", data["sales"]),
            [
                "id",
                "date",
                "customer",
                "product",
                "category",
                "location",
                "quantity",
                "price",
                "unit_cost",
                "sales_value",
                "total_hpp",
            ],
        ),
        ("Recaps", [r for r in data["sales"] if r.get("is_recap")], ["source_sheet", "sku", "product", "category", "location", "period_start", "period_end", "quantity", "sales_value", "total_hpp", "cost_basis"]),
        ("Master", data.get("master_exports", []), ["sku", "upc", "product", "category", "cost_bali", "cost_lombok", "cost_makassar"]),
        ("Rafaksi", data.get("claim_exports", []), ["supplier_code", "supplier_name", "period_start", "period_end", "quantity", "amount"]),
        ("Summary", [{**data["metrics"], "missing_hpp_count": data.get("missing_hpp_count", 0), "note": data.get("note", "")}], list(data["metrics"]) + ["missing_hpp_count", "note"]),
    ]
    for name, rows, headers in groups:
        if name in {"Recaps", "Master", "Rafaksi"} and not rows:
            continue
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append([row.get(key) for key in headers])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 23
    buffer = BytesIO()
    wb.save(buffer)
    wb.close()
    buffer.seek(0)
    return buffer
