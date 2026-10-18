"""Dated sales recaps: preserve source totals; never round them into transactions."""
from datetime import date
from decimal import Decimal, InvalidOperation
from uuid import NAMESPACE_URL, uuid5
import hashlib
import re

from fastapi import HTTPException
from openpyxl.utils.cell import column_index_from_string
from app.models.records import Sale
from app.services.excel import ALIASES, MAX_ROWS, checked_workbook, column_name, detect_header

MONTHS = ["JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI", "JULI", "AGUSTUS", "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER"]
TOTALS = {"total", "grand total", "subtotal", "sub total"}


def describe_report(ws, header_row):
    headers = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    def unique(names):
        found = [i for i, v in enumerate(headers) if v is not None and column_name(v) in names]
        return found[0] if len(found) == 1 else None
    product = unique(ALIASES["product"])
    cost, tax_cost = unique({"cost"}), unique({"cost ppn"})
    dated = []
    for index, value in enumerate(headers):
        match = re.fullmatch(r"QTY[ _]+SALES\s*\((\d{1,2})\s+([A-Z]+)\)", str(value).strip().upper())
        if not match or match[2] not in MONTHS:
            continue
        sales = [i for i, name in enumerate(headers) if str(name).strip().upper() == f"SALES VALUE ({match[1]} {match[2]})"]
        if len(sales) == 1:
            dated.append((index, sales[0], int(match[1]), MONTHS.index(match[2]) + 1))
    if product is None or cost is None or tax_cost is None or len(dated) != 1:
        return None
    quantity, sales, day, month = dated[0]
    if not 1 <= day <= 31:
        return None
    return {
        "kind": "dated_sales_report", "day": day, "month": month,
        "period_label": f"{day} {MONTHS[month - 1].title()}",
        "columns": {"product": product, "quantity": quantity, "sales": sales,
                    "cost": cost, "tax_cost": tax_cost,
                    "region": unique({"region"}), "area": unique({"area"}),
                    "margin": unique({"margin"}), "promo_margin": unique({"margin promo"})},
    }


def amount(cell):
    if cell is None or cell.data_type == "f" or isinstance(cell.value, bool) or not isinstance(cell.value, (int, float, Decimal)):
        raise ValueError("Nilai angka aktual belum tersedia; rumus hanya didukung pada COST+PPN sederhana.")
    try:
        value = Decimal(str(cell.value))
    except InvalidOperation:
        raise ValueError("Nilai angka tidak valid.")
    if not value.is_finite() or value < 0 or value > Decimal("1e18"):
        raise ValueError("Nilai jumlah, penjualan, atau HPP negatif/tidak valid.")
    return value


def cost_value(row, column, columns, number):
    cell = row[column]
    if cell.data_type != "f":
        return amount(cell)
    # Evaluate only a same-row COST times a literal multiplier from the workbook.
    # No arbitrary formula evaluation, external links, cached values or tax assumptions.
    match = re.fullmatch(r"=\$?([A-Z]{1,3})\$?(\d+)\*(\d+(?:\.\d+)?)", str(cell.value).replace(" ", "").upper())
    if column != columns["tax_cost"] or not match or int(match[2]) != number or column_index_from_string(match[1]) - 1 != columns["cost"]:
        raise ValueError("Rumus HPP belum dapat diperiksa. Gunakan salinan berisi nilai HPP aktual.")
    value = amount(row[columns["cost"]]) * Decimal(match[3])
    if value > Decimal("1e12"):
        raise ValueError("HPP melewati batas nilai transaksi yang didukung.")
    return value


def margin_basis(row, columns, number):
    bases = set()
    for name in ["margin", "promo_margin"]:
        index = columns[name]
        if index is None:
            continue
        value = str(row[index].value).replace(" ", "").upper()
        match = re.fullmatch(r"=1-\$?([A-Z]{1,3})\$?(\d+)/\$?([A-Z]{1,3})\$?(\d+)", value)
        if match and int(match[2]) == number and int(match[4]) == number:
            cost_column = column_index_from_string(match[1]) - 1
            if cost_column in [columns["cost"], columns["tax_cost"]]:
                bases.add(cost_column)
    if len(bases) != 1:
        raise ValueError("Dasar HPP pada rumus margin belum jelas atau berbeda dalam satu baris. Periksa COST dan COST+PPN pada sheet ini.")
    return bases.pop()


def prepare_report(content, sheet, year, owner="local"):
    if type(year) is not int or not 1 <= year <= 9999:
        raise HTTPException(400, "Isi tahun laporan yang benar.")
    wb = checked_workbook(content)
    try:
        if sheet not in wb.sheetnames:
            raise HTTPException(400, "Sheet tidak ditemukan.")
        ws = wb[sheet]
        header = detect_header(list(ws.iter_rows(max_row=60, values_only=True)))
        profile = describe_report(ws, header) if header else None
        if not profile:
            raise HTTPException(400, "Sheet ini belum didukung sebagai laporan penjualan bertanggal. Pilih sheet detail Double Date.")
        try:
            day = date(year, profile["month"], profile["day"]).isoformat()
        except ValueError:
            raise HTTPException(400, "Tanggal pada laporan tidak sesuai dengan tahun yang dipilih.")
        columns = profile["columns"]
        records, samples, issues, basis_counts = [], [], [], {}
        zero_rows = total_rows = count = 0
        sales_total, hpp_total = Decimal(0), Decimal(0)
        hpp_complete = sales_complete = True
        signature = hashlib.sha256(content).hexdigest() + f":{sheet}:{year}:dated-report-v1"
        for number, row in enumerate(ws.iter_rows(min_row=header + 1), header + 1):
            if number - header > MAX_ROWS:
                raise HTTPException(400, "Maksimal 20.000 baris per impor.")
            populated = [cell.value for cell in row if cell.value is not None]
            if not populated:
                continue
            product = row[columns["product"]].value
            if (isinstance(product, str) and product.strip().casefold() in TOTALS) or (product is None and str(populated[0]).strip().casefold() in TOTALS):
                total_rows += 1
                continue
            try:
                if row[columns["product"]].data_type == "f" or not isinstance(product, str) or not product.strip():
                    raise ValueError("Nama produk aktual belum tersedia.")
                quantity = amount(row[columns["quantity"]])
                sales = amount(row[columns["sales"]])
                if quantity == 0 and sales == 0:
                    zero_rows += 1
                    continue
                if quantity < 1 or quantity != quantity.to_integral_value() or quantity > 1000000:
                    raise ValueError("Jumlah unit harus bilangan bulat positif.")
            except ValueError as error:
                issues.append({"row": number, "message": str(error)})
                hpp_complete = False
                sales_complete = False
                continue
            count += 1
            sales_total += sales
            cost = hpp = basis = None
            try:
                cost_column = margin_basis(row, columns, number)
                basis = "COST+PPN" if cost_column == columns["tax_cost"] else "COST"
                cost = cost_value(row, cost_column, columns, number)
                hpp = quantity * cost
                hpp_total += hpp
                basis_counts[basis] = basis_counts.get(basis, 0) + 1
            except ValueError as error:
                issues.append({"row": number, "message": str(error)})
                hpp_complete = False
            location = " / ".join(str(row[columns[key]].value).strip() for key in ["region", "area"] if columns[key] is not None and row[columns[key]].value is not None)
            if any(columns[key] is not None and row[columns[key]].data_type == "f" for key in ["region", "area"]):
                issues.append({"row": number, "message": "Nama lokasi berisi rumus; gunakan nilai lokasi aktual."})
            if len(samples) < 10:
                samples.append({"source_row": number, "date": day, "product": product.strip(), "location": location or None,
                                "quantity": int(quantity), "sales_value": str(sales), "unit_cost": str(cost) if cost is not None else None,
                                "total_hpp": str(hpp) if hpp is not None else None, "cost_basis": basis})
            if cost is None:
                continue
            price = (sales / quantity).quantize(Decimal("0.01"))
            if price * quantity != sales:
                issues.append({"row": number, "message": "Total penjualan tidak dapat diwakili harga satuan dua desimal secara tepat. Penyimpanan ditahan agar total tidak berubah."})
                continue
            if cost != cost.quantize(Decimal("0.01")):
                issues.append({"row": number, "message": "HPP sumber memiliki lebih dari dua desimal. Penyimpanan saat ini mendukung dua desimal; nilai sumber tidak dibulatkan."})
                continue
            try:
                record = Sale.model_validate({"date": day, "product": product, "quantity": int(quantity), "price": price,
                                              "unit_cost": cost, "location": location or None}).model_dump(mode="json")
            except Exception:
                issues.append({"row": number, "message": "Nilai produk, lokasi, harga, atau HPP melewati batas transaksi yang didukung."})
                continue
            record["id"] = str(uuid5(NAMESPACE_URL, f"{owner}:{signature}:{number}"))
            records.append(record)
        if not count:
            issues.append({"row": None, "message": "Tidak ada penjualan dengan jumlah unit positif pada sheet ini."})
        view = {
            "kind": "dated_sales_report", "sheet": sheet, "date": day, "count": count,
            "can_import": bool(count and not issues and len(records) == count),
            "sample": samples, "excluded_zero_rows": zero_rows, "excluded_total_rows": total_rows,
            "cost_basis": basis_counts, "issue_count": len(issues), "issues": issues[:10],
            "totals": {"sales": str(sales_total) if sales_complete else None, "hpp": str(hpp_total) if hpp_complete else None,
                       "gross": str(sales_total - hpp_total) if hpp_complete else None},
            "notice": "Ini rekap penjualan per produk/lokasi, bukan daftar transaksi individual. Tanggal memakai hari/bulan pada laporan dan tahun yang Anda pilih. Total penjualan berasal dari SALES VALUE; HPP mengikuti rumus margin masing-masing baris. Biaya operasional tidak diimpor. Data belum disimpan.",
        }
        return view, records
    finally:
        wb.close()
