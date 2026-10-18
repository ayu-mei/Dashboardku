"""Dashboard lokal: jalankan dengan python app.py."""
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from threading import RLock
from uuid import uuid4
import os
import tempfile

from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook, load_workbook

app = Flask(__name__)
app.config['DATA_FILE'] = Path(__file__).parent / 'data' / 'refraksi.xlsx'
lock = RLock()
HEADERS = {
    'Sales': ['id', 'date', 'customer', 'product', 'quantity', 'price', 'unit_cost'],
    'Expenses': ['id', 'date', 'description', 'amount'],
}


def money(value):
    return Decimal(str(value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)


def read_data():
    path = Path(app.config['DATA_FILE'])
    if not path.exists():
        return {name: [] for name in HEADERS}
    wb = load_workbook(path, data_only=False)
    try:
        result = {}
        for name, headers in HEADERS.items():
            if name not in wb or list(next(wb[name].values)) != headers:
                raise ValueError('Struktur Excel tidak sesuai. Gunakan file yang dibuat aplikasi.')
            result[name] = [dict(zip(headers, row)) for row in list(wb[name].values)[1:] if row[0]]
        return result
    finally:
        wb.close()


def write_data(data):
    path = Path(app.config['DATA_FILE'])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for name, headers in HEADERS.items():
        ws = wb.create_sheet(name)
        ws.append(headers)
        ws.freeze_panes = 'A2'
        for record in data[name]:
            ws.append([record[h] for h in headers])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 23
    fd, temporary = tempfile.mkstemp(suffix='.xlsx', dir=path.parent)
    os.close(fd)
    try:
        wb.save(temporary)
        os.replace(temporary, path)
    finally:
        wb.close()
        if os.path.exists(temporary):
            os.unlink(temporary)


def validate(kind, payload):
    result = {}
    try:
        result['date'] = date.fromisoformat(payload.get('date', '')).isoformat()
    except (ValueError, TypeError):
        raise ValueError('Tanggal wajib diisi dengan benar.')
    for key in (['customer', 'product'] if kind == 'Sales' else ['description']):
        value = payload.get(key)
        if not isinstance(value, str) or not value.strip() or len(value.strip()) > 150:
            raise ValueError('Nama/deskripsi wajib diisi, maksimal 150 karakter.')
        result[key] = value.strip()
    for key in (['quantity', 'price', 'unit_cost'] if kind == 'Sales' else ['amount']):
        try:
            value = Decimal(str(payload.get(key, '')))
            if not value.is_finite() or value < 0 or value > Decimal('1000000000000'):
                raise ValueError()
            if key == 'quantity':
                if value < 1 or value != value.to_integral_value() or value > 1000000:
                    raise ValueError()
                result[key] = int(value)
            else:
                if value != money(value):
                    raise ValueError()
                result[key] = float(value)
        except (InvalidOperation, ValueError):
            raise ValueError('Jumlah harus bilangan bulat positif; uang tidak negatif, maksimal 2 desimal.')
    return result


def summarize(data, year):
    sales = [r for r in data['Sales'] if r['date'].startswith(str(year) + '-')]
    expenses = [r for r in data['Expenses'] if r['date'].startswith(str(year) + '-')]
    total = sum((money(r['price']) * r['quantity'] for r in sales), Decimal(0))
    cost = sum((money(r['unit_cost']) * r['quantity'] for r in sales), Decimal(0))
    operational = sum((money(r['amount']) for r in expenses), Decimal(0))
    gross = total - cost
    net = gross - operational
    monthly = [sum((money(r['price']) * r['quantity'] for r in sales if int(r['date'][5:7]) == m), Decimal(0)) for m in range(1, 13)]
    return dict(sales=sales, expenses=expenses, metrics=dict(sales=float(total), hpp=float(cost), gross=float(gross), gross_margin=float(gross / total * 100) if total else None, operational=float(operational), net=float(net), net_margin=float(net / total * 100) if total else None), monthly=list(map(float, monthly)))


@app.get('/')
def index():
    return render_template('dashboard.html')


@app.get('/api/dashboard')
def dashboard():
    try:
        year = int(request.args.get('year', date.today().year))
        if not 1 <= year <= 9999:
            raise ValueError('Tahun tidak valid.')
        with lock:
            return jsonify(summarize(read_data(), year))
    except (ValueError, TypeError) as exc:
        return jsonify(error=str(exc)), 400
    except (OSError, KeyError):
        return jsonify(error='Excel tidak dapat dibaca. Tutup Excel dan periksa file data.'), 503


@app.route('/api/<kind>', methods=['POST'])
@app.route('/api/<kind>/<record_id>', methods=['PUT', 'DELETE'])
def mutate(kind, record_id=None):
    sheet = {'sales': 'Sales', 'expenses': 'Expenses'}.get(kind)
    if not sheet:
        return jsonify(error='Data tidak ditemukan.'), 404
    try:
        with lock:
            data = read_data()
            existing = next((r for r in data[sheet] if r['id'] == record_id), None)
            if request.method != 'POST' and existing is None:
                return jsonify(error='Data tidak ditemukan. Muat ulang dashboard.'), 404
            if request.method == 'DELETE':
                data[sheet].remove(existing)
            else:
                payload = request.get_json(silent=True)
                if not isinstance(payload, dict):
                    raise ValueError('Data formulir tidak valid.')
                record = validate(sheet, payload)
                record['id'] = record_id or str(uuid4())
                if existing:
                    data[sheet].remove(existing)
                data[sheet].append(record)
            write_data(data)
        return jsonify(ok=True)
    except ValueError as exc:
        return jsonify(error=str(exc)), 400
    except OSError:
        return jsonify(error='Gagal menyimpan. Tutup file refraksi.xlsx di Excel lalu coba lagi. Data belum diubah.'), 503


@app.get('/download')
def download():
    from io import BytesIO
    with lock:
        path = Path(app.config['DATA_FILE'])
        if not path.exists():
            write_data(read_data())
        content = BytesIO(path.read_bytes())
    return send_file(content, as_attachment=True, download_name='refraksi.xlsx')


if __name__ == '__main__':
    with lock:
        if not Path(app.config['DATA_FILE']).exists():
            write_data(read_data())
    app.run(host='127.0.0.1', port=5000, debug=False)
